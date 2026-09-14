from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = 'Criterios_Revision'

headers = [
    'DOI', 'Fuente', 'Titulo', 'Anio', 'Evaluador',
    'Decision_Titulo', 'Decision_Resumen', 'Motivo_Exclusion',
    'Criterio_Inclusion', 'Observaciones'
]
ws.append(headers)

header_fill = PatternFill('solid', fgColor='1F4E78')
header_font = Font(color='FFFFFF', bold=True)
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center')

# Example rows
rows = [
    ['10.1001/jama.2023.14530', 'PUBMED', 'Single-Dose Psilocybin Treatment for Major Depressive Disorder: A Randomized Clinical Trial.', '2023', '', 'Incluir', 'Incluir', '', 'Tema central: psilocybin + depresion; diseño clínico', ''],
    ['10.1056/NEJMoa2206443', 'PUBMED', 'Single-Dose Psilocybin for a Treatment-Resistant Episode of Major Depression.', '2022', '', 'Incluir', 'Incluir', '', 'Tema central: psilocybin + depression resistente', ''],
]
for r in rows:
    ws.append(r)

# Validation dropdowns
valid_decision = DataValidation(type='list', formula1='"Incluir,Revisar,Excluir"', allow_blank=False)
valid_decision.prompt = 'Selecciona: Incluir, Revisar o Excluir'
valid_decision.error = 'Debes elegir Incluir, Revisar o Excluir'
valid_decision.add('F2:F1000')
valid_decision.add('G2:G1000')
ws.add_data_validation(valid_decision)

motivo = DataValidation(type='list', formula1='"Titulo no relacionado,Resumen no relevante,Erratum o comentario,Duplicado,Otro"', allow_blank=True)
motivo.prompt = 'Motivo de exclusión o revisión'
motivo.error = 'Elige un motivo entre las opciones previstas'
motivo.add('H2:H1000')
ws.add_data_validation(motivo)

# Freeze panes and auto filter
ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions

# Wrap text and fit widths
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='top')

widths = [20, 15, 75, 12, 20, 16, 16, 35, 45, 30]
for idx, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(idx)].width = w

# Second sheet with criteria guidance
ws2 = wb.create_sheet('Guia_Criterios')
ws2.append(['Seccion', 'Descripcion'])
criteria = [
    ['Criterio de inclusion de titulo', 'Titulo menciona psilocybin, psilocin, Psilocybe, o tratamiento con psilocibina.'],
    ['Criterio de inclusion de abstract', 'Resumen describe estudio o revision sobre psilocybin/psilocin y una pregunta clinica, preclinica, farmacologica o de biosintesis.'],
    ['Criterio de exclusion de titulo', 'Titulo sin relacion con psilocybin, o es noticia/comentario/erratum sin estudio.'],
    ['Criterio de exclusion de abstract', 'Resumen no ofrece datos relevantes, es un comentario/editorial o no es de investigacion.'],
    ['Regla de decision', 'Incluir si titulo y resumen cumplen inclusion; Revisar si hay duda; Excluir si el titulo o el abstract no aportan evidencia relevante.'],
]
for row in criteria:
    ws2.append(row)

for cell in ws2[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center')

for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='top')

ws2.freeze_panes = 'A2'
ws2.auto_filter.ref = ws2.dimensions
for idx in range(1, 3):
    ws2.column_dimensions[get_column_letter(idx)].width = 80

# Save workbook
out = 'criterios_inclusion_rechazo_titulo_abstract.xlsx'
wb.save(out)
print(out)
